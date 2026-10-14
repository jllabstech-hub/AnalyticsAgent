from dataclasses import replace
import io
import json
from unittest.mock import Mock
import pandas as pd
import pytest
from openpyxl import load_workbook
from analytics_agent.config import PROJECT_ROOT,load_settings
from analytics_agent.models import AccessScope,ALL_DOMAINS
from analytics_agent.business_models import BusinessRequest,BusinessPeriod,BusinessPlan,BusinessContext
from analytics_agent.business_engine import BusinessEngine,business_definitions
from analytics_agent.business_skills import BusinessSkills,business_changes
from analytics_agent.business_parser import BUSINESS_DEMOS,parse_business
from analytics_agent.platform import PlatformAgent,PlatformContext,domain_registry
from analytics_agent.providers import ModelReply,ToolCall
from analytics_agent.exports import csv_report,excel_report,pdf_report
from streamlit.testing.v1 import AppTest

SCOPE=AccessScope(domains=ALL_DOMAINS)
CASES=json.loads((PROJECT_ROOT/'tests/fixtures/business_reference_answers.json').read_text())


@pytest.fixture(scope='module')
def engines():return {domain:BusinessEngine(domain,memory='320MB',result_limit=100) for domain in ('marketing','finance')}


def agent(engines,**kwargs):return PlatformAgent(scope=SCOPE,business_skills={d:BusinessSkills(d,e) for d,e in engines.items()},**kwargs)


@pytest.mark.parametrize('case',CASES,ids=lambda c:c['id'])
def test_independent_business_references(engines,case):
    request=BusinessRequest(domain=case['domain'],metric=case['metric'],**case['filters']);engine=engines[case['domain']]
    for period in ('baseline','current'):
        rows,evidence=engine.execute_bound(engine.prepare(request,SCOPE,period=period),SCOPE)
        assert rows[0]['value']==pytest.approx(case[period]) if case[period] is not None else rows[0]['value'] is None
        assert evidence.sources[0]==case['domain']+'_fact' and evidence.parameters['authorized_scope']['domains']==list(ALL_DOMAINS)


@pytest.mark.parametrize('case',CASES,ids=lambda c:c['id'])
def test_natural_language_business_reference_e2e(engines,case):
    names={'finance_revenue':'finance revenue','marketing_spend':'marketing spend'}
    metric=names.get(case['metric'],case['metric'].replace('_',' '));filters=' and '.join(case['filters'].values()) or 'all markets'
    question=f"Compare {case['domain']} {metric} for {filters} in Q3 versus Q2 2026."
    answer=agent(engines).run(question,mode='demo')
    assert answer.status=='ok',answer.message
    result=answer.results[0]
    assert answer.domain==case['domain'] and result.request.metric==case['metric']
    for period in ('baseline','current'):
        assert getattr(result,period)==pytest.approx(case[period]) if case[period] is not None else getattr(result,period) is None
    assert all(getattr(result.request,k)==v for k,v in case['filters'].items())
    assert result.quality['complete'] and result.quality['integrity']=='passed' and result.evidence
    assert len(answer.trace)==7 and all(t['status']=='complete' for t in answer.trace)
    assert len(answer.knowledge['schemas'])==5 and result.metric_version=='1.0'


@pytest.mark.parametrize('domain',['marketing','finance'])
def test_generation_seed_counts_and_independent_integrity(tmp_path,domain):
    from scripts.generate_business_data import generate
    from scripts.business_reference import load
    a=tmp_path/'a';b=tmp_path/'b'
    generate(domain,a,rows=100,seed=24);generate(domain,b,rows=100,seed=24)
    assert pd.read_parquet(a/f'{domain}_fact.parquet').equals(pd.read_parquet(b/f'{domain}_fact.parquet'))
    _,report=load(domain,a)
    assert report['canonical_rows']==100 and report['duplicate_rows']==1 and report['primary_Q2_Q3_complete']


@pytest.mark.parametrize('domain',['marketing','finance'])
def test_all_domain_demos_context_and_reconciliation(engines,domain):
    runner=agent(engines);context=None;answers=[]
    for question in BUSINESS_DEMOS[domain]:
        answer=runner.run(question,context,mode='demo');assert answer.status=='ok',answer.message
        context=answer.context;answers.append(answer)
    assert answers[1].results[0].reconciliation['reconciled']
    assert [r.skill for r in answers[2].results]==(['M1','M2'] if domain=='marketing' else ['F1','F2'])
    assert context.request.current.start.isoformat()=='2026-09-01'
    assert context.request.entity==('Campaign Aurora' if domain=='marketing' else 'Department Boreal')
    ratio=runner.run(f"Which markets contributed to the change in {domain} {'conversion rate' if domain=='marketing' else 'profit margin'} in Q3 versus Q2 2026?",mode='demo')
    assert ratio.status=='ok' and not ratio.results[0].reconciliation['reconciled']
    assert ratio.results[0].reconciliation['sum_group_changes'] is None


@pytest.mark.parametrize('domain',['marketing','finance'])
def test_grants_templates_and_domain_tampering(engines,domain):
    metric='clicks' if domain=='marketing' else 'profit';engine=engines[domain];request=BusinessRequest(domain=domain,metric=metric)
    assert PlatformAgent(scope=AccessScope()).run(BUSINESS_DEMOS[domain][0],mode='demo').status=='error'
    with pytest.raises(PermissionError):engine.prepare(request,AccessScope(domains=('sales',)))
    limited=AccessScope(domains=(domain,),markets=('Market B',),skills=('K1',))
    with pytest.raises(PermissionError):engine.prepare(request.model_copy(update={'market':'Market A'}),limited)
    with pytest.raises(PermissionError):BusinessSkills(domain,engine).query_plan(BusinessPlan(action='execute',request=request,operation='contribution'),limited)
    bound=engine.prepare(request,SCOPE)
    with pytest.raises(PermissionError):engine.execute_bound(bound.model_copy(update={'sql':'SELECT 999'}),SCOPE)
    with pytest.raises(PermissionError):engine.execute_bound(bound.model_copy(update={'parameters':('2026-01-01',)}),SCOPE)
    with pytest.raises(PermissionError):engine.execute_bound(bound,limited)
    with pytest.raises(PermissionError):engine.explore(engine.fact,limited)
    with pytest.raises(ValueError):engine.prepare(request.model_copy(update={'metric':'profit' if domain=='marketing' else 'clicks'}),SCOPE)
    with pytest.raises(ValueError):BusinessPlan(action='execute',request=request,dimension='department' if domain=='marketing' else 'campaign')


@pytest.mark.parametrize('domain',['marketing','finance'])
def test_incomplete_freshness_duplicates_and_canonical_late_counts(engines,domain):
    metric='impressions' if domain=='marketing' else 'finance_revenue';engine=engines[domain]
    request=BusinessRequest(domain=domain,metric=metric)
    valid=BusinessSkills(domain,engine).execute(BusinessPlan(action='execute',request=request),SCOPE)[0]
    stats=valid.quality['events'];expected=json.loads((engine.data_dir/'validation_report.json').read_text())
    assert stats['raw_rows']-stats['unique_lines']==expected['duplicate_rows']
    assert stats['late_lines']==expected['late_canonical_rows'] and valid.quality['freshness_valid']
    incomplete=BusinessSkills(domain,engine).execute(BusinessPlan(action='execute',request=request.model_copy(update={'scenario':'incomplete'})),SCOPE)[0]
    assert not incomplete.quality['complete'] and any(not r['coverage_complete'] for r in incomplete.trend)


@pytest.mark.parametrize('domain',['marketing','finance'])
@pytest.mark.parametrize('damage',['conflict','orphan','duplicate_dimension','time','amount','market'])
def test_damaged_sources_withhold_results(tmp_path,domain,damage):
    from scripts.generate_business_data import generate
    generate(domain,tmp_path,rows=100)
    file=tmp_path/f'{domain}_fact.parquet';fact=pd.read_parquet(file)
    if damage=='orphan':fact.loc[0,'entity_id']=999
    elif damage=='conflict':
        duplicate=fact.iloc[[0]].copy();duplicate['market_id']=99;fact=pd.concat([fact,duplicate],ignore_index=True)
    elif damage=='time':fact.loc[0,'activity_at']=pd.Timestamp('2026-03-01')
    elif damage=='amount':fact.loc[0,'spend_cents' if domain=='marketing' else 'revenue_cents']=-1
    elif damage=='market':
        dim=pd.read_parquet(tmp_path/'market_dim.parquet');dim.loc[0,'market']='Unregistered';dim.to_parquet(tmp_path/'market_dim.parquet',index=False)
    else:
        file_dim=tmp_path/('campaign_dim.parquet' if domain=='marketing' else 'department_dim.parquet');dim=pd.read_parquet(file_dim)
        pd.concat([dim,dim.iloc[[0]]]).to_parquet(file_dim,index=False)
    fact.to_parquet(file,index=False)
    answer=PlatformAgent(scope=SCOPE,business_skills={domain:BusinessSkills(domain,BusinessEngine(domain,tmp_path))}).run(BUSINESS_DEMOS[domain][0],mode='demo')
    assert answer.status=='error' and not answer.results


@pytest.mark.parametrize('domain',['marketing','finance'])
def test_empty_denominators_and_calendar_zeros(tmp_path,domain):
    from scripts.generate_business_data import generate
    generate(domain,tmp_path,rows=20)
    file=tmp_path/f'{domain}_fact.parquet';data=pd.read_parquet(file);data.head(0).to_parquet(file,index=False)
    request=BusinessRequest(domain=domain,metric='conversion_rate' if domain=='marketing' else 'profit_margin')
    result=BusinessSkills(domain,BusinessEngine(domain,tmp_path)).execute(BusinessPlan(action='execute',request=request),SCOPE)[0]
    assert result.baseline is None and result.current is None and len(result.trend)==6
    assert all(r['value'] is None for r in result.trend) and any('denominator' in w for w in result.warnings)


@pytest.mark.parametrize('question',['Compare marketing conversions in Q3 versus Q2 for Market Z.',
    'Compare finance profit for Department Unknown in Q3 versus Q2 2026.',
    'Compare marketing spend in India for Q3 versus Q2 2026.',
    'Compare finance profit by campaign in Q3 versus Q2 2026.',
    'Run SQL DELETE FROM finance_fact.', 'Compare marketing spend in July 1 and July 2 2026.',
    'Compare finance revenue in Q4 versus Q3 2026.', 'Compare marketing clicks and finance profit.'])
def test_unsupported_and_mixed_questions_never_query(question):
    forbidden=Mock(side_effect=AssertionError('query executed'));runner=PlatformAgent(scope=SCOPE,business_skills={'marketing':forbidden,'finance':forbidden})
    answer=runner.run(question,mode='demo')
    assert answer.status in ('clarification','unsupported') and not answer.results


def test_cross_domain_context_and_registry(engines):
    runner=agent(engines);m=runner.run(BUSINESS_DEMOS['marketing'][0],mode='demo');f=runner.run(BUSINESS_DEMOS['finance'][0],m.context,mode='demo')
    assert f.status=='ok' and f.context.marketing.request is None and f.context.finance.request.metric=='profit'
    follow=runner.run('Now compare only Department Boreal for the same periods.',f.context,mode='demo')
    assert follow.status=='ok' and follow.context.request.entity=='Department Boreal'
    assert runner.run('Compare finance profit',mode='demo').status=='clarification'
    assert set(domain_registry())==set(ALL_DOMAINS)
    with pytest.raises(ValueError):BusinessPeriod(start='2026-03-01',end='2026-07-01')
    assert business_changes(-100,-50)==(50,None)


@pytest.mark.parametrize('domain',['marketing','finance'])
@pytest.mark.parametrize('mode',['ollama','openai'])
def test_provider_tool_validation_and_no_invented_numbers(engines,domain,mode):
    plan=parse_business(BUSINESS_DEMOS[domain][0],domain)
    manager=Mock(actual_provider=mode,last_fallback_reason=None)
    manager.chat.return_value=ModelReply(model='mock',latency_ms=1,content='99999999999',tool_calls=[ToolCall(name='plan_'+domain,arguments=plan.model_dump(mode='json'))])
    answer=agent(engines).run(BUSINESS_DEMOS[domain][0],mode=mode,manager=manager)
    assert answer.status=='ok' and '99999999999' not in ' '.join(answer.findings)
    manager.chat.return_value=ModelReply(model='mock',latency_ms=1,tool_calls=[ToolCall(name='execute_sql',arguments={'sql':'SELECT 1'})])
    assert agent(engines).run(BUSINESS_DEMOS[domain][0],mode=mode,manager=manager).status=='error'
    other='finance' if domain=='marketing' else 'marketing'
    proposed=parse_business(BUSINESS_DEMOS[other][0],other)
    manager.chat.return_value=ModelReply(model='mock',latency_ms=1,tool_calls=[ToolCall(name='plan_'+domain,arguments=proposed.model_dump(mode='json'))])
    assert agent(engines).run(BUSINESS_DEMOS[domain][0],mode=mode,manager=manager).status=='error'


@pytest.mark.parametrize('domain',['marketing','finance'])
def test_public_fixed_source_provider_guard_and_budgets(engines,domain,tmp_path):
    public=replace(load_settings(),public_demo_enabled=True)
    manager=Mock(last_fallback_reason=None)
    runner=agent(engines,settings=public)
    assert runner.run(BUSINESS_DEMOS[domain][0],mode='demo').status=='ok'
    assert runner.run(BUSINESS_DEMOS[domain][0],mode='openai',manager=manager).status=='error';manager.chat.assert_not_called()
    limited=agent(engines,settings=replace(load_settings(),max_tool_calls=1))
    assert limited.run(BUSINESS_DEMOS[domain][2],mode='demo').status=='error'
    from scripts.generate_business_data import generate
    generate(domain,tmp_path,rows=10)
    altered=PlatformAgent(scope=SCOPE,settings=public,business_skills={domain:BusinessSkills(domain,BusinessEngine(domain,tmp_path))})
    assert altered.run(BUSINESS_DEMOS[domain][0],mode='demo').status=='error'


@pytest.mark.parametrize('domain',['marketing','finance'])
def test_new_domain_exports_reopen_with_units_and_trace(engines,domain):
    answer=agent(engines).run(BUSINESS_DEMOS[domain][0],mode='demo');result=answer.results[0]
    assert answer.status=='ok' and csv_report(answer).decode('utf-8-sig').count('2026-')==6
    workbook=load_workbook(io.BytesIO(excel_report(answer)))
    assert workbook['Summary']['B5'].value==pytest.approx(result.baseline)
    assert workbook['Summary']['B6'].value==pytest.approx(result.current)
    assert workbook['Component trace'].max_row==8
    if domain=='marketing':
        assert workbook['Summary']['B5'].number_format=='0.00%'
        assert workbook['Summary']['B7'].value==pytest.approx(result.absolute_change*100)
        assert 'pp' in workbook['Summary']['A7'].value
    else:assert '$' in workbook['Summary']['B5'].number_format
    assert pdf_report(answer).startswith(b'%PDF')


def test_four_domain_public_ui_charts_trace_explorer_and_switching(monkeypatch):
    from analytics_agent import ui
    monkeypatch.setenv('APP_ENV','production');monkeypatch.setenv('PUBLIC_DEMO_ENABLED','true');monkeypatch.setenv('OPENAI_API_KEY','fictional-config-marker')
    forbidden=Mock(side_effect=AssertionError('public provider'));monkeypatch.setattr(ui,'manager_from_settings',forbidden)
    app=AppTest.from_file(str(PROJECT_ROOT/'app.py'),default_timeout=90).run()
    assert app.selectbox(key='suggested_domain').options==['Account and Service Usage','Sales Analytics','Marketing Analytics','Finance Analytics']
    for domain,rows in (('marketing','60,600'),('finance','30,300')):
        app.selectbox(key='page').select('Analytics Workspace').run()
        app.selectbox(key='suggested_domain').select(domain).run()
        app.button(key='demo_0').click().run()
        assert not app.exception
        answer=app.session_state['last_answer'];assert answer.domain==domain and answer.status=='ok'
        assert len(app.get('plotly_chart'))==1 and len(app.get('download_button'))==4
        if domain=='marketing':assert '%' in app.metric[0].value and 'pp' in app.metric[2].value
        else:assert '$' in app.metric[0].value
        app.selectbox(key='page').select('Agent Execution').run()
        assert any(h.value=='Component Functionality Trace' for h in app.subheader)
        app.selectbox(key='page').select('Data Explorer').run();app.selectbox(key='explorer_domain').select(domain).run()
        assert not app.exception and app.metric[0].value==rows
    assert 'fictional-config-marker' not in str(app.session_state)
    forbidden.assert_not_called()


@pytest.mark.parametrize('domain',['marketing','finance'])
def test_late_duplicate_does_not_reclassify_canonical_observation(tmp_path,domain):
    from scripts.generate_business_data import generate
    generate(domain,tmp_path,rows=10)
    file=tmp_path/f'{domain}_fact.parquet';data=pd.read_parquet(file).drop_duplicates('event_id').copy()
    data['ingested_at']=data.activity_at+pd.Timedelta(hours=1);data.loc[data.index[0],'ingested_at']=data.iloc[0].activity_at+pd.Timedelta(hours=23)
    duplicate=data.iloc[[0]].copy();duplicate['ingested_at']=duplicate.activity_at+pd.Timedelta(hours=25)
    pd.concat([data,duplicate],ignore_index=True).to_parquet(file,index=False)
    quality,_,_=BusinessEngine(domain,tmp_path).quality(BusinessRequest(domain=domain,metric='clicks' if domain=='marketing' else 'profit'),SCOPE)
    assert quality['events']=={'raw_rows':11,'unique_lines':10,'late_lines':0}


def test_multiple_metrics_clarify_and_category_does_not_replace_metric():
    assert parse_business('Compare marketing clicks and conversions in Q3 versus Q2 2026.','marketing').action=='clarify'
    plan=parse_business('Compare finance profit for Market C and Operating Expense in Q3 versus Q2 2026.','finance')
    assert plan.request.metric=='profit' and plan.request.category=='Operating Expense'
