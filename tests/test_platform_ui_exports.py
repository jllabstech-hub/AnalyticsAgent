import io
import json
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from analytics_agent.config import PROJECT_ROOT
from analytics_agent.models import AccessScope
from analytics_agent.platform import PlatformAgent
from analytics_agent.sales_parser import SALES_DEMOS
from analytics_agent.exports import csv_report,excel_report,pdf_report


def test_sales_exports_preserve_query_values_and_component_trace():
    answer=PlatformAgent(scope=AccessScope(domains=("sales",))).run(SALES_DEMOS[0],mode="demo")
    csv=csv_report(answer).decode("utf-8-sig")
    assert "2026-04" in csv and "units_sold" in csv
    workbook=load_workbook(io.BytesIO(excel_report(answer)),data_only=True)
    assert workbook.sheetnames==["Summary","Results","Component trace"]
    assert workbook["Summary"]["B5"].value==1748539.73 and workbook["Summary"]["B6"].value==3218017.11
    assert workbook["Component trace"].max_row==8 and workbook["Results"].max_row==7
    assert pdf_report(answer).startswith(b"%PDF")
    failed=answer.model_copy(update={"status":"error"})
    for builder in (csv_report,excel_report,pdf_report):
        with pytest.raises(ValueError): builder(failed)


def test_public_sales_ui_queries_charts_exports_and_trace(monkeypatch):
    from analytics_agent import ui
    from unittest.mock import Mock
    monkeypatch.setenv("APP_ENV","production")
    monkeypatch.setenv("PUBLIC_DEMO_ENABLED","true")
    monkeypatch.setenv("OPENAI_API_KEY", "fictional-private-marker")
    forbidden=Mock(side_effect=AssertionError("Public sales attempted provider access"))
    monkeypatch.setattr(ui,"manager_from_settings",forbidden)
    app=AppTest.from_file(str(PROJECT_ROOT/"app.py"),default_timeout=60).run()
    app.selectbox(key="suggested_domain").select("sales").run()
    app.button(key="demo_0").click().run()
    assert not app.exception
    answer=app.session_state["last_answer"]
    assert answer.domain=="sales" and answer.status=="ok"
    assert app.metric[0].value=="$1,748,539.73" and app.metric[1].value=="$3,218,017.11"
    assert len(answer.trace)==7 and all(e["status"]=="complete" for e in answer.trace)
    assert len(app.get("plotly_chart"))==1 and len(app.get("download_button"))==4
    app.selectbox(key="page").select("Agent Execution").run()
    assert any(header.value=="Component Functionality Trace" for header in app.subheader)
    app.selectbox(key="page").select("Data Explorer").run()
    app.selectbox(key="explorer_domain").select("sales").run()
    assert not app.exception and app.metric[0].value=="50,500"
    assert "fictional-private-marker" not in str(app.session_state)
    forbidden.assert_not_called()
