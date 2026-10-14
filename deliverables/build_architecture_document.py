"""Create the team architecture handover using the bundled document runtime."""
from pathlib import Path
import json
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables'
QA = OUT / 'architecture_qa'
OUT.mkdir(exist_ok=True)
QA.mkdir(exist_ok=True)
DOCX = OUT / 'BDC_Analytics_Agent_Architecture_and_Technical_Handover_V1.docx'
evaluation = json.loads((ROOT / 'reports/evaluation.json').read_text(encoding='utf-8'))
performance = json.loads((ROOT / 'reports/performance.json').read_text(encoding='utf-8'))
integrity = json.loads((ROOT / 'data/synthetic/validation_report.json').read_text(encoding='utf-8'))

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin, section.bottom_margin = Inches(.66), Inches(.65)
section.left_margin, section.right_margin = Inches(.72), Inches(.72)
section.header_distance, section.footer_distance = Inches(.25), Inches(.3)
WIDTH = 6.83
styles = doc.styles
for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption', 'List Bullet', 'List Number'):
    style = styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:ascii'), 'Arial')
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'), 'Arial')
    borders = style.element.xpath('.//w:pBdr')
    for border in borders:
        border.getparent().remove(border)
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.line_spacing = 1.08
styles['Normal'].paragraph_format.widow_control = True
styles['Title'].font.size = Pt(25)
styles['Title'].font.bold = True
styles['Title'].paragraph_format.space_after = Pt(12)
styles['Subtitle'].font.size = Pt(12)
styles['Subtitle'].paragraph_format.space_after = Pt(10)
styles['Heading 1'].font.size = Pt(17)
styles['Heading 1'].paragraph_format.space_before = Pt(0)
styles['Heading 1'].paragraph_format.space_after = Pt(10)
styles['Heading 2'].font.size = Pt(12)
styles['Heading 2'].paragraph_format.space_before = Pt(11)
styles['Heading 2'].paragraph_format.space_after = Pt(5)
styles['Caption'].font.size = Pt(9)
styles['Caption'].font.italic = False
styles['Caption'].paragraph_format.space_after = Pt(8)
for name in ('List Bullet', 'List Number'):
    styles[name].font.size = Pt(10.5)
    styles[name].paragraph_format.space_after = Pt(4)
    styles[name].paragraph_format.line_spacing = 1.06

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run('BDC Analytics Agent V1  |  ')
run.font.name, run.font.size = 'Arial', Pt(8)
run.font.color.rgb = RGBColor(0, 0, 0)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.core_properties.title = 'BDC Analytics Agent Architecture and Technical Handover'
doc.core_properties.subject = 'Prototype V1 architecture data contracts controls and operations'
doc.core_properties.author = 'BDC Analytics Agent Project'
doc.core_properties.keywords = 'synthetic analytics, DuckDB, Streamlit, Ollama, OpenAI'
doc.core_properties.comments = ''

def p(text, bold_lead=None):
    paragraph = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        paragraph.add_run(bold_lead).bold = True
        paragraph.add_run(text[len(bold_lead):])
    else:
        paragraph.add_run(text)
    return paragraph

def h(text, level=2):
    return doc.add_heading(text, level)

def page(title):
    doc.add_page_break()
    h(title, 1)

def bullet(text):
    return doc.add_paragraph(text, style='List Bullet')

def code(text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.0
    for idx, line in enumerate(text.splitlines()):
        run = paragraph.add_run(('\n' if idx else '') + line)
        run.font.name, run.font.size = 'Consolas', Pt(8.5)
    return paragraph

def table(headers, rows, widths, compact=False, centered=()):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        for attr, value in (('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')):
            e.set(qn(f'w:{attr}'), value)
        borders.append(e)
    props.append(borders)
    for i, values in enumerate([headers] + list(rows)):
        row = t.rows[0] if i == 0 else t.add_row()
        trpr = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit')
        trpr.append(no_split)
        if i == 0:
            repeat = OxmlElement('w:tblHeader')
            trpr.append(repeat)
        for j, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge, val in [('top','75'), ('bottom','75'), ('left','95'), ('right','95')]:
                e = OxmlElement('w:' + edge)
                e.set(qn('w:w'), val)
                e.set(qn('w:type'), 'dxa')
                margins.append(e)
            tcpr.append(margins)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '183B56' if i == 0 else ('F3F6FA' if i % 2 else 'FFFFFF'))
            tcpr.append(shade)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.04
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in centered else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(str(value))
            run.font.name = 'Arial'
            run.font.size = Pt(9 if compact else 9.5)
            run.font.bold = i == 0
            run.font.color.rgb = RGBColor(255,255,255) if i == 0 else RGBColor(0,0,0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def diagram(path, boxes, arrows, size=(1800, 870)):
    image = Image.new('RGB', size, 'white')
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 31)
    bold = ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf', 34)
    navy, teal = '#183B56', '#007F86'
    for start, end in arrows:
        draw.line((start, end), fill=teal, width=5)
        x, y = end
        sx, sy = start
        if abs(x-sx) > abs(y-sy):
            d = 1 if x > sx else -1
            draw.polygon([(x,y), (x-d*17,y-10), (x-d*17,y+10)], fill=teal)
        else:
            d = 1 if y > sy else -1
            draw.polygon([(x,y), (x-10,y-d*17), (x+10,y-d*17)], fill=teal)
    for rect, title, detail in boxes:
        draw.rounded_rectangle(rect, radius=13, fill='#F3F6FA', outline=navy, width=3)
        x1,y1,x2,y2=rect
        lines=[(title,bold)]+[(line,font) for line in detail.split('\n')]
        height=len(lines)*42
        y=y1+(y2-y1-height)/2
        for line, face in lines:
            bounds=draw.textbbox((0,0),line,font=face)
            x=x1+((x2-x1)-(bounds[2]-bounds[0]))/2
            draw.text((x,y),line,font=face,fill=navy)
            y+=42
    image.save(path)

diagram(QA/'architecture.png', [
    ((25,45,555,180),'Streamlit workspace','Chat charts settings and evidence'),
    ((635,45,1165,180),'Analytics agent','Validated and bounded plan'),
    ((1245,45,1775,180),'Provider manager','Ollama or selected OpenAI'),
    ((635,295,1165,430),'Analytical skills','K1 comparison K2 groups K3 cohorts'),
    ((25,295,555,430),'Semantic definitions','Versioned YAML metrics'),
    ((1245,295,1775,430),'Quality and evidence','Coverage populations reconciliation'),
    ((635,545,1165,680),'DuckDB query engine','Approved templates and safe views'),
    ((635,760,1165,865),'Synthetic Parquet','Six source datasets')
], [((555,112),(635,112)),((1165,112),(1245,112)),((900,180),(900,295)),
    ((555,362),(635,362)),((900,430),(900,545)),((1165,362),(1245,362)),
    ((900,680),(900,760))])

diagram(QA/'data_relationships.png', [
    ((25,15,555,150),'Market reference','Market and UTC timezone'),
    ((635,15,1165,150),'Users','Identity registration and activation'),
    ((1245,15,1775,150),'Quality metadata','Dataset scenario and period'),
    ((25,265,555,400),'Service membership','User service and eligibility interval'),
    ((635,265,1165,400),'Device mapping','User device and validity interval'),
    ((1245,265,1775,400),'Service activity','Event user service and device')
], [((555,82),(635,82)),((900,150),(290,265)),
    ((900,150),(900,265)),((900,150),(1510,265))], size=(1800,415))

# Page 1
doc.add_paragraph('BDC Analytics Agent', style='Title')
doc.add_paragraph('Architecture and Technical Handover', style='Title')
doc.add_paragraph('Prototype V1', style='Subtitle')
p('Prepared for the engineering analytics and leadership teams\n14 October 2026')
p('This document describes the implemented application, its technical contracts, and the controls that keep analytical answers traceable. It gives the team the information needed to operate the local prototype, review its calculations, extend approved capabilities, and prepare a controlled hosted release.')
p('Current assessment. The synthetic prototype passes Windows and GitHub Linux CI, including a separately verified real OpenAI workflow. Production approval requires broader live model evaluation, real identity-provider lifecycle testing, and hosted reliability validation.', 'Current assessment.')
table(['Item','Implemented scope'],[
    ('Data','100,000 fictional accounts; July to September 2026; six Parquet datasets'),
    ('Analytics','Six governed metrics; K1 comparison, K2 contribution, K3 cohorts'),
    ('Interaction','Streamlit; default Ollama; optional explicitly selected OpenAI; explicit demo mode'),
    ('Recorded verification',f"{evaluation['total']} tests: {evaluation['passed']} passed, {evaluation['failed']} failed, {evaluation['skipped']} skipped"),
    ('Release status','Windows and Linux CI pass; broader live integration and hosting pending')
], [1.55,5.28])
h('Reading guide')
p('Pages 2 to 5 explain the architecture, technology stack, and warehouse. Pages 6 to 9 explain metrics, quality controls, contracts, orchestration, and providers. Pages 10 to 13 cover security, execution resources, the interface, and deployment. Pages 14 to 16 cover measured validation, development stages, and team handover.')
p('All people, account identifiers, markets, services, and device models in the warehouse are fictional. No enterprise data is connected. The application does not establish confirmed business root causes from these synthetic observations.')

# Page 2
page('1 System architecture')
p('A single Python application separates presentation, planning, calculation, and validation. The UI calls one bounded AnalyticsAgent. The agent resolves a typed request and executes registered skills. Skills call approved DuckDB functions and return values, warnings, chart records, and query evidence.')
picture = doc.add_paragraph().add_run().add_picture(str(QA/'architecture.png'), width=Inches(WIDTH))
picture._inline.docPr.set('descr','Component diagram showing Streamlit, the agent, optional providers, YAML definitions, analytical skills, the controlled DuckDB engine, Parquet data, and quality evidence.')
doc.add_paragraph('Figure 1  Logical components and approved calculation path', style='Caption')
table(['Boundary','Responsibility'],[
    ('Presentation','Render computed results and public execution stages; store nonsecret session preferences.'),
    ('Planning','Resolve metric, periods, filters, intent, and approved skill tasks; clarify ambiguity.'),
    ('Semantics and skills','Apply versioned metric definitions and reusable analytical methods.'),
    ('Queries and source data','Execute parameterized templates over bundled, time-valid Parquet views.'),
    ('Validation and evidence','Check coverage, populations, denominators, duplicates, and reconciliation; attach sources.')
], [1.7,5.13], compact=True)
h('Architectural approach')
p('The backend does not depend on Streamlit for calculations. This permits command-line tests and later frontend replacement. V1 uses no vector database, distributed worker system, agent framework, or microservices. New complexity should follow a measured requirement such as a larger warehouse or sustained concurrency.')
p('Model output is treated as untrusted planning input. The model never provides SQL, executable code, permissions, or numerical evidence. Demo, Ollama, and OpenAI modes share the same skills and calculation engine.')
p('Recognized questions carry approved period and skill contracts. Original model scopes are checked before correction, identical tasks are coalesced, and each chart has a unique interface key.')

# Page 3
page('2 Technology stack and source modules')
p('The recorded build was tested on Windows and GitHub Linux CI with Python 3.14.5. Direct dependencies are pinned in requirements.txt. This is not a complete transitive lock; target installations must be checked and scanned again.')
table(['Technology','Pinned version','Purpose'],[
    ('Streamlit','1.65.0','Workspace, session state, UI tests, OIDC integration'),
    ('DuckDB','1.5.6','Embedded analytical SQL over Parquet'),
    ('pandas / NumPy','3.0.6 / 2.4.6','Synthetic generation and independent references'),
    ('PyArrow / Faker','24.0.0 / 40.43.0','Parquet serialization and fictional labels'),
    ('Plotly','7.1.0','Interactive trend and group charts'),
    ('Pydantic / PyYAML','2.14.1 / 6.0.3','Validated contracts and metric registry'),
    ('OpenAI SDK / requests','3.29.0 / 2.34.2','Cloud adapter and Ollama HTTP adapter'),
    ('Authlib / python dotenv','1.8.0 / 1.2.4','OIDC support and local environment loading'),
    ('pytest','9.1.1','Automated analytical, provider, and interface tests')
], [1.95,1.25,3.63], compact=True)
table(['Source location','Responsibility'],[
    ('app.py and analytics_agent/ui.py','Entry point, chat, Settings, Explorer, Evaluation'),
    ('config.py and models.py','Settings and immutable public contracts'),
    ('metrics.yaml and semantics.py','Business definitions and registry loading'),
    ('engine.py and skills.py','Warehouse access, approved queries, K1 to K3'),
    ('agent.py and providers.py','Planning, context, findings, adapters, fallback'),
    ('access.py','Production identity claims and server authorization map'),
    ('scripts and tests','Generation, references, diagnostics, evaluation, regression'),
    ('render.yaml and .github/workflows','Prepared hosting and Linux quality workflow')
], [2.7,4.13], compact=True)
p('Additional quality tooling is declared in requirements-dev.txt, including pip-audit 2.10.1. The project installer was upgraded to pip 26.2.1 after the dependency audit found installer advisories.')

# Page 4
page('3 Synthetic warehouse and event model')
p('Generation uses seed 20260701. Markets A to E, Services Alpha to Gamma, and devices Orion One, Orion Plus, Nova Mini, and Nova Fold provide approved dimensions. Registration begins before the reporting window, allowing new and existing cohorts.')
table(['Dataset','Rows','Grain'],[
    ('users',f"{integrity['row_counts']['users']:,}",'One fictional user and account'),
    ('service_activity',f"{integrity['row_counts']['service_activity']:,}",'One usage event row; event_id can repeat'),
    ('service_membership',f"{integrity['row_counts']['service_membership']:,}",'User and service eligibility interval'),
    ('device_mapping',f"{integrity['row_counts']['device_mapping']:,}",'User and device association interval'),
    ('market_reference','5','One fictional market'),
    ('data_quality_status','30','Scenario, source dataset, reporting month')
], [2.0,.95,3.88], compact=True, centered=(1,))
h('Users schema')
table(['Fields','Stored type','Meaning'],[
    ('user_id','int64','Primary user key'),
    ('account_id and display_name','large_string','Fictional account ID and display label'),
    ('market','large_string','Market A to E; reference join'),
    ('registered_at and activated_at','timestamp[us]','UTC instants; activation may be null')
], [2.45,1.3,3.08], compact=True)
h('Activity schema')
table(['Fields','Stored type','Meaning'],[
    ('event_id','large_string','Stable event identity for deduplication'),
    ('user_id','int64','Foreign key to users'),
    ('service and device_id','large_string','Service and device used for the event'),
    ('event_at and ingested_at','timestamp[ns]','Usage time and pipeline arrival time'),
    ('event_type','large_string','Synthetic usage event type')
], [2.45,1.3,3.08], compact=True)
p('The primary warehouse contains 1,180,343 unique events, 5,902 duplicate rows, and 23,429 unique events arriving more than 24 hours late. The complete primary snapshot includes those late events. These are measured source counts, not expected answers embedded in the agent.')

# Page 5
page('4 Association schemas and temporal joins')
table(['Dataset','Fields and stored types'],[
    ('service_membership','user_id int64; service large_string; eligible_from and eligible_to timestamp[us]; membership_cohort large_string'),
    ('device_mapping','user_id int64; device_id and device_model large_string; valid_from and valid_to timestamp[us]'),
    ('market_reference','market and timezone large_string; is_synthetic bool'),
    ('data_quality_status','scenario, dataset, status large_string; period_start, period_end, available_through, loaded_at, snapshot_as_of timestamp[us]; complete bool')
], [1.85,4.98])
picture = doc.add_paragraph().add_run().add_picture(str(QA/'data_relationships.png'), width=Inches(WIDTH))
picture._inline.docPr.set('descr','Data relationship diagram showing market reference linked to users, users linked to membership, device mappings and activity, and separate dataset scenario coverage metadata.')
doc.add_paragraph('Figure 2  User associations and source coverage metadata', style='Caption')
h('Relationships and reporting boundaries')
bullet('users.user_id joins activity, membership, and device associations. users.market joins market_reference.market.')
bullet('Activity joins membership on user_id and service, with eligible_from <= event_at < eligible_to.')
bullet('Activity joins device history on user_id and device_id, with valid_from <= event_at < valid_to. Activity must occur on or after activation.')
bullet('User keys must be unique. Membership intervals for one user and service must not overlap. Device intervals for one user must not overlap. Orphan references and invalid intervals fail quality validation.')
p('All timestamps are stored without timezone offsets and interpreted exclusively as UTC. Reporting ranges are half-open: August is [2026-08-01, 2026-09-01), and September is [2026-09-01, 2026-10-01). Valid requests must remain inside July to September 2026.')
h('Separate incomplete scenario')
p('data/synthetic/incomplete/service_activity.parquet omits September 24 through 30. Other sources remain unchanged. Quality metadata identifies 23 available September days. The primary August and September partitions remain complete and are used for the executive comparison.')
h('Historical freshness')
p('Data-quality records use loaded_at 4 October 2026 and snapshot_as_of 5 October 2026. Freshness is evaluated against that declared historical snapshot. It is not live pipeline freshness or a measurement against the current date.')

# Page 6
page('5 Business metrics and calculation rules')
p('Each YAML definition records metric ID, version, business meaning, calculation, source datasets, required joins, timezone, grain, supported dimensions, numerator, denominator, owner, and quality requirements. All six metrics support approved market, service, and device filters.')
table(['Metric','Calculation and population'],[
    ('MAU\nmau 1.0','Distinct activated eligible users with an event in the selected period. Users active in multiple selected services are counted once.'),
    ('Mean DAU\ndau 1.0','Sum of daily distinct active users divided by all calendar days in the period, including covered days with no activity.'),
    ('New Users\nnew_users 1.0','Distinct eligible users registered within the selected period. Registration and service eligibility define the population.'),
    ('Active Existing Users\nactive_existing_users 1.0','Distinct active eligible users registered before the selected period starts.'),
    ('Activation Rate\nactivation_rate 1.0','New eligible users activated within the same reporting period divided by all new eligible registrations in that period.'),
    ('Service Adoption Rate\nservice_adoption_rate 1.1','Distinct active eligible users divided by users eligible during the period. Device and membership associations must actually intersect within the selected period.')
], [2.0,4.83])
h('Population and device rules')
p('Activity metrics use the device valid at the event time. New-user and activation populations use the device association valid immediately before period end. Adoption uses period-overlapping device and membership associations in both numerator and denominator; this prevents a device switch from retaining an active numerator while removing its denominator.')
h('Comparisons and units')
p('Absolute change is current minus baseline. Relative percentage change is absolute change divided by baseline times 100; a zero baseline makes relative change undefined. Ratio deltas are displayed as percentage points. For example, 40% to 45% is +5 percentage points and +12.5% relative change.')
p('A zero denominator produces an undefined rate, not zero. Ratios are not added across groups. New and existing cohort membership is defined separately at each period start and does not represent a fixed-cohort retention study.')

# Page 7
page('6 Reusable skills and validation')
table(['Skill','Inputs','Output and safeguards'],[
    ('K1 Metric Comparison','Metric, two periods, approved filters','Baseline, current, absolute and relative change; full daily supporting DAU trend.'),
    ('K2 Contribution Breakdown','K1 request plus market, service, or device dimension','Group values and changes; contribution percentages only when total reconciliation is valid.'),
    ('K3 Cohort Comparison','K1 request plus new_existing, registration, or membership','Period-relative new/existing groups, registration months, or Standard/Premium membership cohorts.')
], [1.5,2.0,3.33])
p('Skill version is 1.1. Every skill executes independently through SkillRegistry and returns a structured SkillResult. No skill needs an LLM. Shared validation covers metric existence, supported dimensions, authorized scope, source coverage, snapshot freshness, duplicates, finite values, denominator validity, joins, and reconciliation.')
h('Data quality behavior')
bullet('Repeated event IDs with identical core payloads are deduplicated, retaining the earliest ingestion row. Conflicting payloads for the same ID withhold findings.')
bullet('Impossible ingestion timestamps, invalid or overlapping association intervals, orphan references, and join cardinality violations withhold findings.')
bullet('Incomplete reporting coverage produces provisional results with visible warnings. Trend days without verified coverage are null gaps; covered zero-activity days are zero.')
bullet('Late events and removed duplicate counts appear in evidence-backed warnings. They do not automatically invalidate the complete historical primary comparison.')
h('When contributions reconcile')
p('K2 compares the sum of group populations with the total in both periods and requires an additive metric. Market partitions can reconcile. Service populations and historical device groups may overlap, so their distinct counts cannot automatically explain an additive total. Contribution shares may exceed 100% or be negative when growth offsets a larger decline.')
p('For the recorded MAU comparison, the net change is -733 users. Market A contributes -4,506 and the other markets contribute +3,773 in aggregate. These observed contributions reconcile; they do not establish a confirmed cause of lower engagement.')

# Page 8
page('7 Agent workflow and typed contracts')
p('The public workflow is Understanding, Skill Selection, Query Execution, Validation, and Findings. The execution trace exposes completed stages and skill names, without hidden model reasoning.')
table(['Contract','Main fields'],[
    ('AnalyticsRequest','metric; baseline/current Period; optional market, service, device_model; primary/incomplete scenario'),
    ('Plan and Task','execute/clarify/unsupported action; validated request; registered skill tasks; intent'),
    ('AccessScope','Server-controlled approved markets, services, and skills'),
    ('SkillResult','Versions; values; records; trend; warnings; quality; query evidence; reconciliation; duration'),
    ('QueryEvidence','Query ID; template; actual parameters and executed period; scope; sources; version 1.1; latency; rows; cache flag'),
    ('AgentAnswer','Status; provider/model; fallback reason; computed findings; results; follow-ups; trace; context; elapsed time')
], [1.55,5.28], compact=True)
h('One investigation')
bullet('Resolve an approved metric, two reporting periods, filters, and intent from the question and last authorized request.')
bullet('In LLM mode require exactly one plan_investigation tool call. Validate every field using frozen Pydantic models with extra fields forbidden.')
bullet('Recheck population permissions and skill allowlists. Enforce configured task count, depth, time, and result limits.')
bullet('Execute K1, K2, and K3 as appropriate. Validate calculations and generate findings from returned values.')
bullet('Optionally let the model select and order computed observation IDs. Keep the primary change, completeness warning, and causal limitation visible.')
bullet('Return authorized context and suggested follow-ups that retain the selected metric, periods, and scope.')
p('Deterministic demo mode uses a bounded parser for documented questions. Unknown filters, unsupported data, and ambiguous dates are refused or clarified rather than silently changing the population. It is explicitly enabled and does not pretend to be a model response.')
h('Programmatic skill execution')
p('This local fictional-data example runs the default August and September MAU comparison. Production callers must supply the scope derived from authenticated identity.')
example = '''from analytics_agent.engine import Engine
from analytics_agent.models import AccessScope, AnalyticsRequest
from analytics_agent.skills import SkillRegistry, ToolArguments
registry = SkillRegistry(Engine())
result = registry.execute(
    "K1", ToolArguments(request=AnalyticsRequest()), AccessScope()
)
print(result.model_dump_json(indent=2))'''
code(example)
(QA/'skill_example.py').write_text(example + '\n', encoding='utf-8')

# Page 9
page('8 LLM providers and fallback behavior')
p('Provider defines health and chat operations. OllamaProvider uses REST; OpenAIProvider uses the OpenAI SDK Chat Completions interface with approved tools and structured output. ProviderManager records the actual provider and model used for every answer.')
table(['Condition','Behavior'],[
    ('Ollama selected','Use configured Ollama only. Failure never selects OpenAI.'),
    ('OpenAI selected and configured','Use the explicitly selected OpenAI model.'),
    ('OpenAI key missing or recoverable timeout, connection, service failure','If fallback is enabled, check Ollama reachability and configured model availability; then attempt local inference.'),
    ('Authentication, permission, invalid request, or invalid model output','Return a safe error; do not fall back.'),
    ('Both providers unavailable','Return an error and offer explicit deterministic demo mode.'),
    ('Successful fallback within an answer','Keep the remaining answer local and expose the fallback reason.')
], [2.8,4.03], compact=True)
h('Ollama adapter')
p('Default origin is http://localhost:11434 and model qwen3:8b. qwen3:1.7b is a configurable smaller option. The adapter supports /api/tags discovery, /api/show metadata, /api/chat, JSON schemas, and tool calls. It limits response size, rejects redirects, classifies errors, and validates malformed discovery/tool responses. GPU diagnostics detect hardware where possible; detection does not establish actual inference offload.')
h('OpenAI adapter')
p('Default model is gpt-4.1-mini and can be changed to a compatible tool-capable model. Automatic SDK retries are disabled. Timeouts and error classifications are explicit. Passive health reports configuration without a billed inference call. Test Connection makes an actual request and reports inference verification only when a response succeeds.')
h('Endpoint and credential boundaries')
p('Secrets remain server-side. Remote Ollama requires an HTTPS origin and server bearer token for a protected proxy; the proxy must enforce authentication. Render cannot access Ollama on the Windows laptop through localhost. No model downloads or software installation are initiated by the application.')
p('Recorded verification. OpenAI Test Connection and one complete MAU comparison have succeeded. Ollama installation is awaiting approval; its inference and broader live model behavior remain unverified.', 'Recorded verification.')

# Page 10
page('9 Authentication authorization and security')
p('Local mode provides the full fictional demonstration scope. Production mode requires Streamlit OIDC sign-in and an explicit server-owned authorization map. Render forces production mode. Missing configuration stops the app before analytics or provider controls appear.')
h('Identity and scope checks')
bullet('Require HTTPS redirect and metadata URLs, nonempty client credentials, and a cookie secret of at least 32 characters. Placeholder configuration is rejected.')
bullet('Read identity claims from st.user after Streamlit verification. Require issuer iss, subject sub, and a finite unexpired exp value.')
bullet('Resolve the exact issuer and subject key in the access map. Each entry must specify nonempty approved markets, services, and skills.')
bullet('Apply scope in the agent, skill registry, and SQL filters, including requests that do not specify a market or service.')
bullet('Clear retained answers and conversation when identity or effective scope changes. Data Explorer requires full market and service access.')
h('Execution and secret controls')
p('Public tool contracts contain no SQL, code, filesystem paths, or permission overrides. Only internal approved query templates run. DuckDB external access is restricted to bundled source files; extensions are disabled and configuration is locked. Application queries use SELECT, WITH, or DESCRIBE through the controlled layer.')
p('OPENAI_API_KEY and OLLAMA_API_TOKEN are excluded from configuration representations, browser controls, logs, and shared caches. Local .env and actual Streamlit secret files are ignored. Provider errors are redacted, and internally owned HTTP clients are closed after use.')
h('Production sign in configuration')
p('Use .streamlit/secrets.example.toml as the configuration template. Register the exact HTTPS /oauth2callback address with the approved identity provider. On Render provide secrets.toml through Secret Files at /etc/secrets/secrets.toml. Store account grants under the exact ISSUER|SUBJECT key with explicit approved values.')
p('Identity authentication and analytical authorization are separate controls. The local suite tests simulated claims, denied access, scoped UI behavior, and scope revocation. Actual signed login/logout, provider claim availability, expiry, and account revocation require live IdP testing. A known-vulnerability scan and source credential scan are not a penetration test or security certification.')

# Page 11
page('10 Execution resources configuration and caching')
p('Streamlit caches the Engine resource, not a live DuckDB connection. Each operation creates and closes its own in-memory database connection. A shared reentrant warehouse lock bounds database execution across sessions, coalesces identical cold cache requests, and respects queue deadlines.')
table(['Setting','Default','Meaning'],[
    ('LLM_PROVIDER','ollama','Explicit provider selection'),
    ('OLLAMA_BASE_URL','http://localhost:11434','Local origin; remote HTTPS requires token'),
    ('OLLAMA_MODEL / OPENAI_MODEL','qwen3:8b / gpt-4.1-mini','Configurable compatible models'),
    ('OPENAI_FALLBACK_TO_OLLAMA','true','Eligible failures only'),
    ('LLM_TIMEOUT_SECONDS','30','Provider request budget'),
    ('MAX_TOOL_CALLS / INVESTIGATION_DEPTH','6 / 3','Bounded plan; both limits apply'),
    ('INVESTIGATION_TIMEOUT_SECONDS','120','Total investigation deadline'),
    ('RESULT_LIMIT','500','Query and analytical record limit'),
    ('DUCKDB_MEMORY_LIMIT','512MB','DuckDB tracked memory budget'),
    ('Engine query timeout / threads','15 seconds / 1','Interrupted execution; bounded CPU'),
    ('APP_ENV','local','production requires configured access')
], [2.65,1.6,2.58], compact=True)
h('Cache correctness')
p('The engine retains at most 32 approved results. Keys include request, periods, operation, authorization scope, source file sizes and modification times, and resource settings. Returned values are deep copies. Authorization runs on cache hits, and evidence marks hits. Quality and DAU trend results can be reused across metrics with request-correct evidence.')
h('Memory and restart behavior')
p('Disk spill is disabled. DuckDB memory is only part of total memory; Python, Streamlit, charts, and model HTTP clients add overhead. Cached results are process-local and rebuild after restart. Data files are bundled and immutable during normal execution, so persistent writable cloud storage is not needed.')
p('Environment variables take precedence over local .env values. Invalid or nonfinite limits are rejected. Backend file watching is disabled because changed configuration classes can leave stale imports in a running process. Perform a clean server restart after backend or secret configuration changes.')

# Page 12
page('11 Application interface and demonstration outputs')
table(['Section','User-facing behavior'],[
    ('Analytics Workspace','Chat and follow-ups alongside computed KPI cards, DAU trend, group changes, cohorts, and findings.'),
    ('Agent Execution','Public workflow stages, registered skills, completion status, and elapsed time.'),
    ('Settings','Provider and model selectors, fallback toggle, connection test, actual provider and last fallback reason.'),
    ('Data Explorer','Approved read-only schemas, row counts, and fictional samples; scope controlled.'),
    ('Evaluation Dashboard','Actual aggregate evaluation report, supported skills, measurements, and limitations.')
], [1.8,5.03])
p('The interface uses a white background, navy and teal accents, responsive KPI wrapping, and Plotly charts. Synthetic data and explicit demo mode are labeled. Incomplete coverage appears above KPI cards and in the chat response. Supporting trend charts are explicitly daily active users even when another metric is selected.')
h('Five demonstration workflows')
bullet('Compare August and September active users.')
bullet('Which markets contributed most to the decline? Retain context from the first question.')
bullet('Investigate the September decline across markets and cohorts.')
bullet('Can we trust the September activity data?')
bullet('Now compare only Market A for the same periods. Retain the authorized comparison context.')
h('Recorded primary comparison')
table(['Population','August 2026','September 2026','Change'],[
    ('All markets MAU','66,025','65,292','-733'),
    ('Market A','13,149','8,643','-4,506'),
    ('Market B','13,132','14,960','+1,828'),
    ('Active existing users','58,798','58,314','-484'),
    ('Active new users','7,227','6,978','-249')
], [2.0,1.5,1.7,1.63], compact=True, centered=(1,2,3))
p('These values are validation examples from executed calculations, not routing rules or embedded findings. Clear the conversation before an independent incomplete scenario test; that September MAU is 62,411 and must be marked provisional.')

# Page 13
page('12 Local operation and Render preparation')
p('Run the bundled data locally without an API key by explicitly enabling Deterministic demo mode in the sidebar. Use the project virtual environment directly to avoid PowerShell activation-policy issues.')
code("Set-Location 'D:\\Office\\AnalyticsAgent'\n.\\.venv\\Scripts\\python.exe -m streamlit run app.py\n# Open http://127.0.0.1:8501")
p('For a fresh machine, create .venv, upgrade pip to 26.2.1, install requirements.txt, run pip check, and copy .env.example to ignored .env. Configure secrets only on the server. Parquet datasets are already bundled; generation is needed only to reproduce or intentionally rebuild them.')
h('Prepared hosting contract')
table(['Item','Prepared value'],[
    ('Service','Render Python Web Service; auto-deploy disabled'),
    ('Runtime and resources','Python 3.14.5; paid 1c-2g plan with 2 GB RAM'),
    ('Health endpoint','/_stcore/health'),
    ('State','Bundled Parquet plus ephemeral in-memory DuckDB; no persistent disk'),
    ('Access','APP_ENV=production and configured OIDC plus account scopes')
], [1.7,5.13], compact=True)
code('python -m pip install --upgrade pip==26.2.1\npython -m pip install --prefer-binary -r requirements.txt')
p('Render start command', bold_lead='Render start command')
code('streamlit run app.py --server.address 0.0.0.0 --server.port $PORT --server.headless true')
p('Bind to the assigned PORT. Configure OPENAI_API_KEY in environment variables or use a separately secured reachable Ollama endpoint. Localhost on Render refers to Render, not the Windows laptop.')
h('Release procedure')
p('Review and push the source and synthetic datasets to the authorized jllabstech-hub/AnalyticsAgent GitHub repository. Run the prepared Linux quality workflow. After hosting and resource authorization, configure staging secrets and account grants, deploy, verify sign-in and all five demos, restart, and repeat the comparison. Monitor failures, memory, latency, and API expenditure before release approval.')
p('The 2 GB plan is prepared, not purchased. A 256 MB DuckDB budget failed warehouse join validation; a 512 MB total host should not be assumed sufficient. Source is uploaded to main and Linux CI passes; public application hosting remains unauthorized.')

# Page 14
page('13 Validation performance and reproducibility')
p('The full suite compares production calculations against independent pandas references and a hand-calculated small warehouse. Reference answers are stored under tests/fixtures and are never imported by the agent. Provider tests use controlled mocks; live tests require explicit opt-in.')
table(['Recorded check','Result'],[
    ('Full test suite',f"{evaluation['passed']} passed; {evaluation['failed']} failed; {evaluation['skipped']} skipped; {evaluation['total']} total"),
    ('Independent analytical comparisons','30 of 30 matched'),
    ('Defined natural-language workflows','30 of 30 matched in deterministic mode'),
    ('Executive demonstrations','Five executed successfully'),
    ('Startup and restart','Two local starts returned HTTP 200 and ok'),
    ('Dependency advisory scan','83 packages; zero known records; no skipped packages'),
    ('Cold / warm investigation',f"{performance['end_to_end_ms']/1000:.3f} seconds / {performance['warm_end_to_end_ms']:.3f} milliseconds"),
    ('Concurrent investigations','Four of four succeeded with identical totals'),
    ('Peak backend working set',f"{performance['peak_process_memory_mb']:.2f} MiB on Windows")
], [2.55,4.28], compact=True)
p('Evidence is from the recorded local validation run on 14 October 2026 India time. Accuracy is 100% on the defined synthetic reference suite; it is not enterprise accuracy or live-model routing accuracy. Backend timings exclude live LLM latency and additional hosted/UI overhead. Four simultaneous requests do not establish sustained production capacity.')
p('Separate live evidence confirms OpenAI inference and the full August and September K1 browser comparison: 66,025 to 65,292 users, change -733, completed in 49.12 seconds with demo mode off. See reports/openai_live_validation.json. Linux CI also passes the full suite, credential and dependency scans, and both restarts; see reports/github_validation.json.')
h('Reproduce the checks')
code('.\\.venv\\Scripts\\python.exe -m scripts.validate_data\n.\\.venv\\Scripts\\python.exe -m scripts.evaluate\n.\\.venv\\Scripts\\python.exe -m scripts.run_demos\n.\\.venv\\Scripts\\python.exe -m scripts.benchmark\n.\\.venv\\Scripts\\python.exe -m scripts.security_check\n.\\.venv\\Scripts\\python.exe -m scripts.smoke_deployment --port 8765')
p('To rebuild source data and independent references, run scripts.generate_data followed by scripts.validate_data and scripts.reference_answers before evaluation. This replaces the generated synthetic warehouse intentionally. Use requirements-dev.txt for advisory scanning. Verify hosted CI results in the repository Actions tab.')
p('Coverage includes boundaries, overlapping populations, duplicates, incomplete coverage, zero denominators, unsupported requests, scope changes, context, invalid tool calls, fallback, provider failures, and Streamlit startup and browser output.')

# Page 15
page('14 Development stages and implemented corrections')
table(['Stage','Main deliverable','Status'],[
    ('1 Foundation','Environment, modules, configuration, logs','Complete locally'),
    ('2 Data','Seeded six-table fictional warehouse','Complete locally'),
    ('3 Queries','Controlled DuckDB templates and references','Complete locally'),
    ('4 Semantics','Versioned YAML and K1 to K3 skills','Complete locally'),
    ('5 Ollama','Adapter, diagnostics, structured tools','Awaiting install approval'),
    ('6 Providers','OpenAI selection and local fallback','OpenAI K1 live verified'),
    ('7 Agent','Bounded planning, context, findings','Complete locally'),
    ('8 Interface','Workspace, Settings, Explorer, Evaluation','Complete locally'),
    ('9 Evaluation','Independent calculations and regressions','Complete locally'),
    ('10 Hosting','Render blueprint and restart checks','Linux CI pass; Render pending'),
    ('11 EVP demos','Five executed executive workflows','Complete locally'),
    ('12 Hardening','Integration, audit, performance, reports','Complete locally')
], [.98,3.83,2.02], compact=True)
h('Corrections that matter for the team')
p('The final audit corrected service routing and unknown-filter handling, period-relative cohort selection, and adoption numerator/denominator alignment. It added complete calendar trends with null gaps, explicit percentage-point changes, actual query-period evidence, invalid-source rejection, and compulsory primary findings.')
p('Execution hardening bounded shared DuckDB work, coalesced cache misses, and closed owned provider clients. UI hardening preserved actual fallback metadata, wrapped compact KPI cards, and required clean restarts for backend changes. Production access guards and authorization regression tests were added before hosting preparation was considered ready for staging review.')
p('IMPLEMENTATION_STATUS.md records feasible phase completion and external blockers. BUILD_LOG.md records implementation decisions, executed commands, failures, and fixes. Future changes should update both and preserve the shared numerical path across all modes.')

# Page 16
page('15 Team handover and release requirements')
h('Required verification before production approval')
bullet('Model integration: rotate the exposed server key, evaluate the reference suite and failure paths in each enabled provider, and install a local model only after approval. One real OpenAI K1 workflow is already verified.')
bullet('Identity integration: register the actual OIDC client; verify HTTPS login/logout, issuer/subject/expiry claims, expired and revoked users, and restricted market/service access.')
bullet('Hosted reliability: Linux CI passes. In authorized staging, measure sustained load, memory, restart, network failures, and operational monitoring against agreed acceptance targets.')
bullet('Operating controls: assign ownership for identity grants, model configuration, release review, incident handling, and account-level API expenditure limits. These team responsibilities are proposed, not assigned.')
h('How to extend the application')
p('For a new metric, add the versioned YAML definition, update the approved contract and query implementation, calculate independent reference answers, and add population and boundary tests. For a new skill, register its validated inputs and outputs and update the skill allowlist. For a new provider, implement the common adapter contract without changing the calculation layer or weakening selection and fallback controls.')
p('Real enterprise data requires a separately approved source integration, lineage and quality contracts, access review, and new evaluation data. The current synthetic-only release should remain the baseline for those changes. Broader language support and higher traffic also require new measured acceptance tests.')
h('Repository reference index')
table(['Reference','Use'],[
    ('ARCHITECTURE.md','Component boundaries and analytical semantics'),
    ('README.md and DEPLOYMENT_GUIDE.md','Setup, providers, OIDC, GitHub, and Render instructions'),
    ('FINAL_VALIDATION_REPORT.md and PRODUCTION_READINESS.md','Actual evidence, limitations, and release requirements'),
    ('DEMO_GUIDE.md','Five workflows and expected validation examples'),
    ('analytics_agent/metrics.yaml','Authoritative business definitions and versions'),
    ('data/synthetic/validation_report.json','Source counts and injected quality issues'),
    ('reports/evaluation.json and test_results.xml','Actual automated evaluation results'),
    ('reports/performance.json and deployment_smoke.json','Local measurements and startup evidence'),
    ('IMPLEMENTATION_STATUS.md and BUILD_LOG.md','Phase status and implementation history')
], [3.1,3.73], compact=True)
p('This handover documents the implemented prototype and recorded validation evidence, including one successful live OpenAI comparison. Remaining external verification must be completed before production approval. All data is fictional; public application hosting and paid resource creation remain pending authorization.')

doc.save(DOCX)
report = {
    'document': str(DOCX),
    'explicit_pages': 16,
    'paragraphs': len(doc.paragraphs),
    'tables': len(doc.tables),
    'generation_date': '2026-10-14',
    'evaluation_source_utc': evaluation['generated_at_utc'],
    'title': doc.core_properties.title,
}
(QA/'document_structure.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
print(json.dumps(report, indent=2))
