"""Memory-only exports of executed answers. No queries or model calls here."""
import csv
import io
import json
from html import escape


def result_rows(answer):
    if answer.status!="ok" or not answer.results: raise ValueError("Only successful executed answers can be exported.")
    rows=[]
    for result in answer.results:
        for row in result.trend or result.records:
            rows.append({"skill":result.skill,**row})
    if not rows:
        first=answer.results[0]
        rows=[{"baseline":first.baseline,"current":first.current,"change":first.absolute_change}]
    return rows


def csv_report(answer):
    rows=result_rows(answer)
    columns=list(dict.fromkeys(key for row in rows for key in row))
    buffer=io.StringIO(newline="")
    writer=csv.DictWriter(buffer,fieldnames=columns)
    writer.writeheader(); writer.writerows(rows)
    return buffer.getvalue().encode("utf-8-sig")


def excel_report(answer):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    rows=result_rows(answer); first=answer.results[0]
    workbook=Workbook(); summary=workbook.active; summary.title="Summary"
    summary.append(["BDC Analytics - synthetic report","Value"])
    summary.append(["Domain",getattr(answer,"domain","usage")])
    summary.append(["Metric",first.request.metric]); summary.append(["Provider",answer.provider])
    summary.append(["Baseline",first.baseline]); summary.append(["Current",first.current])
    unit=first.definition.get("unit","usd" if first.request.metric in ("net_sales","avg_order_value") else "count")
    ratio=unit=="ratio"
    summary.append(["Absolute change (pp)" if ratio else "Absolute change",first.absolute_change*100 if ratio and first.absolute_change is not None else first.absolute_change]); summary.append(["Relative percent change",first.percentage_change])
    summary.append(["Baseline UTC",f"{first.request.baseline.start} to {first.request.baseline.end} (exclusive)"])
    summary.append(["Current UTC",f"{first.request.current.start} to {first.request.current.end} (exclusive)"])
    summary.append(["Synthetic only","Observed associations do not establish a confirmed cause."])
    data=workbook.create_sheet("Results")
    columns=list(dict.fromkeys(key for row in rows for key in row)); data.append(columns)
    for row in rows: data.append([row.get(column) for column in columns])
    trace=workbook.create_sheet("Component trace")
    trace.append(["Component","Status","Elapsed milliseconds","Evidence"])
    for event in answer.trace:
        details={key:value for key,value in event.items() if key not in ("stage","status","component_id","elapsed_ms")}
        trace.append([event["stage"],event["status"],event.get("elapsed_ms"),json.dumps(details,ensure_ascii=True,default=str)[:32000]])
    for sheet in workbook:
        sheet.freeze_panes="A2"; sheet.sheet_view.showGridLines=False
        for cell in sheet[1]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="183B56")
        sheet.row_dimensions[1].height=30
        for column in sheet.columns:
            letter=column[0].column_letter
            sheet.column_dimensions[letter].width=min(52,max(18,max(len(str(cell.value or "")) for cell in column)+2))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment=Alignment(vertical="top",wrap_text=sheet is trace or sheet is summary)
                if isinstance(cell.value,str) and cell.value.startswith(("=","+","-","@")):
                    cell.data_type="s"
            if sheet is trace: sheet.row_dimensions[row[0].row].height=48
        sheet.auto_filter.ref=sheet.dimensions
    for row in (5,6,7): summary.cell(row,2).number_format='"$"#,##0.00' if unit=="usd" else '0.00%' if ratio and row!=7 else '0.00" pp"' if ratio else '#,##0.00'
    output=io.BytesIO(); workbook.save(output)
    return output.getvalue()


def pdf_report(answer):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak
    rows=result_rows(answer); first=answer.results[0]; buffer=io.BytesIO()
    ratio=first.definition.get("unit")=="ratio"
    styles=getSampleStyleSheet(); styles["Normal"].fontSize=9; styles["Normal"].leading=13
    def para(text,style="Normal"):
        ascii_text=str(text).replace("→","->").replace("–","-").replace("·","|")
        return Paragraph(escape(ascii_text),styles[style])
    story=[para("BDC Analytics | Synthetic Report","Title"),Spacer(1,12),
        para(f"Domain: {getattr(answer,'domain','usage')} | Metric: {first.request.metric} | Provider: {answer.provider}"),
        para(f"Baseline UTC: {first.request.baseline.start} to {first.request.baseline.end} (exclusive)"),
        para(f"Current UTC: {first.request.current.start} to {first.request.current.end} (exclusive)"),Spacer(1,12)]
    for finding in answer.findings: story.extend([para(finding),Spacer(1,6)])
    story.extend([Spacer(1,10),para("Executed result data","Heading2")])
    columns=list(dict.fromkeys(key for row in rows for key in row))
    # Compact report data; full records and SQL are available in CSV/Excel/trace JSON.
    selected=[key for key in ("month","group","period","value","baseline","current","change","units_sold","orders") if key in columns][:6]
    if not selected: selected=columns[:6]
    table_data=[[para(key.replace("_"," ")) for key in selected]]
    for row in rows[:40]:
        values=[]
        for key in selected:
            value=row.get(key)
            display="Undefined" if value is None else f"{value:.2%}" if ratio and key in ("value","baseline","current") else f"{value*100:+.2f} pp" if ratio and key=="change" else f"{value:,.2f}" if isinstance(value,(int,float)) and not isinstance(value,bool) else value
            values.append(para(display))
        table_data.append(values)
    table=Table(table_data,colWidths=[480/len(selected)]*len(selected),repeatRows=1,hAlign="LEFT")
    table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#E7F3F4")),("VALIGN",(0,0),(-1,-1),"TOP"),
        ("BOTTOMPADDING",(0,0),(-1,-1),7),("TOPPADDING",(0,0),(-1,-1),7),("LINEBELOW",(0,0),(-1,0),.5,colors.HexColor("#007F86"))]))
    story.extend([table,Spacer(1,12),para("Quality and source information","Heading2")])
    for warning in dict.fromkeys(w for r in answer.results for w in r.warnings): story.extend([para(warning),Spacer(1,4)])
    story.append(para("Sources: "+", ".join(dict.fromkeys(source for result in answer.results for evidence in result.evidence for source in evidence.sources))))
    story.extend([PageBreak(),para("Component Functionality Trace","Title"),Spacer(1,14)])
    for event in answer.trace:
        story.append(para(f"{event.get('component_id','')} {event['stage']} | {event['status']}","Heading2"))
        keys={"User Query":"execution_mode","Intent Understanding":"request","Domain & Skill Selection":"registry",
            "Knowledge Retrieval":"source","Approved Query Generation":"policy","Execute & Validate":"executed_skills","Results & Insights":"finding_basis"}
        content=event.get(keys.get(event["stage"],""),"See full validated evidence in the downloadable trace JSON.")
        if event["stage"]=="Domain & Skill Selection" and event.get("registry"):
            item=event["registry"]
            content=f"Registered {item['name']} ({item['skill_id']}, version {item['version']}); operations {', '.join(item['operations'])}. Selected {event.get('operation','usage skill plan')}."
        elif event["stage"]=="Knowledge Retrieval" and event.get("schemas"):
            definition=event["definition"]
            content=f"Retrieved {len(event['schemas'])} table schemas, metric {definition['metric_id']} version {definition['version']}, source mappings and validated join rules. {definition['business_meaning']}"
        elif event["stage"]=="Approved Query Generation" and event.get("queries"):
            content=f"Prepared {len(event['queries'])} read-only queries using approved total/trend/breakdown templates. UTC dates and population filters are bound separately; SQL text cannot be supplied by questions or models."
        elif event["stage"]=="Execute & Validate" and event.get("quality"):
            quality=event["quality"]; stats=quality["events"]
            content=f"Executed {', '.join(event['executed_skills'])}. Integrity {quality['integrity']}; coverage complete: {quality['complete']}; freshness valid: {quality['freshness_valid']}. {stats['raw_rows']:,} scoped rows, {stats['unique_lines']:,} canonical observations, {stats['late_lines']:,} late observations."
        elif event["stage"]=="Results & Insights" and "chart_rows" in event:
            content=f"Produced {event['result_count']} skill result(s) and {event['chart_rows']} chart/table rows from executed values. Follow-ups and CSV, Excel, PDF and complete trace JSON are available."
        if isinstance(content,dict): content=json.dumps(content,ensure_ascii=True,default=str)
        story.append(para(str(content)[:1000])); story.append(Spacer(1,5))
    def footer(canvas,document):
        canvas.setFont("Helvetica",8); canvas.setFillColor(colors.HexColor("#66788A"))
        canvas.drawString(42,24,"Fictional data only | Query-backed observations | No enterprise production certification")
        canvas.drawRightString(553,24,str(document.page))
    SimpleDocTemplate(buffer,pagesize=A4,rightMargin=42,leftMargin=42,topMargin=40,bottomMargin=45).build(story,onFirstPage=footer,onLaterPages=footer)
    return buffer.getvalue()
